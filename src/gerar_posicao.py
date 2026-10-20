#!/usr/bin/env python3
"""Gera o relatório "Posição <empresa>" (padrão Exact BR, mesmo layout da ALS).

Uso:
  python3 src/gerar_posicao.py --hoje 2026-10-20            # busca o Omie e grava dados/omie.json
  python3 src/gerar_posicao.py --hoje 2026-10-20 --offline  # usa o último dados/omie.json

Fontes: Omie (títulos, baixas, razão), Webpilot (entradas/webpilot/faturas-*.xlsx),
Itaú (entradas/itau/boletos-vencidos-*.xlsx) e config/bmg.json (saldos documentados e guias).
Regra: saldo de banco só entra se vier de documento (config.saldos_documentados).
"""
import argparse, base64, collections, datetime as dt, glob, html, json, re, time
from pathlib import Path

import openpyxl

import omie

RAIZ = Path(__file__).resolve().parent.parent
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
DIAS = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]


# ------------------------------------------------------------------ utilidades
def d(s):
    return dt.datetime.strptime(s[:10], "%d/%m/%Y").date() if "/" in s else dt.date.fromisoformat(s[:10])


def br(v, dec=2):
    s = f"{abs(v):,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("−" if v < 0 else "") + s


def mm(v):
    return f"R$ {br(v / 1e6)}<span class=\"sm\"> MM</span>" if abs(v) >= 1e6 else f"R$ {br(v / 1e3, 0)}<span class=\"sm\">K</span>"


def dm(x):
    return x.strftime("%d/%m") if x else "—"


def esc(s):
    return html.escape(str(s or ""))


def dig(s):
    return re.sub(r"\D", "", str(s or ""))


def rps_de(s):
    m = re.search(r"(\d+)", str(s or ""))
    return str(int(m.group(1))) if m else None


# ------------------------------------------------------------------ dados
def buscar_omie(cache):
    def pesq(nat, **flt):
        out, pg = [], 1
        while True:
            p = {"nPagina": pg, "nRegPorPagina": 200, "cNatureza": nat}; p.update(flt)
            r = omie.call("financas/pesquisartitulos", "PesquisarLancamentos", p)
            out += r.get("titulosEncontrados", [])
            if pg >= r.get("nTotPaginas", 1):
                return out
            pg += 1; time.sleep(1.2)
    dados = {
        "R": pesq("R", dDtVencDe="01/01/2025", dDtVencAte="31/12/2027"),
        "P": pesq("P", dDtVencDe="01/01/2026", dDtVencAte="31/12/2027"),
        "L": omie.paginas("financas/contacorrentelancamentos", "ListarLancCC", "listaLancamentos"),
        "CLI": omie.paginas("geral/clientes", "ListarClientes", "clientes_cadastro", pag="pagina", rpp="registros_por_pagina", tot="total_de_paginas"),
    }
    cache.parent.mkdir(exist_ok=True)
    cache.write_text(json.dumps(dados, ensure_ascii=False))
    return dados


def saldo_omie(conta, ate):
    e = omie.call("financas/extrato", "ListarExtrato", {"nCodCC": int(conta), "dPeriodoInicial": ate.replace(day=1).strftime("%d/%m/%Y"), "dPeriodoFinal": ate.strftime("%d/%m/%Y")})
    return e.get("nSaldoAtual")


def ler_webpilot():
    """NFS-e → {cliente, navio, pedido}. Usa todas as exportações em entradas/webpilot/."""
    m = {}
    for arq in sorted(glob.glob(str(RAIZ / "entradas" / "webpilot" / "*.xlsx"))):
        ws = openpyxl.load_workbook(arq, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if r and r[0] == "Número")
        col = {c: i for i, c in enumerate(rows[h])}
        for r in rows[h + 1:]:
            if r and r[col["NFSe"]]:
                m[str(r[col["NFSe"]])] = {"cliente": r[col["Cliente"]], "navio": r[col["Navio"]], "pedido": r[col["Pedido(s)"]]}
    return m


def ler_boletos_banco():
    arqs = sorted(glob.glob(str(RAIZ / "entradas" / "itau" / "boletos-vencidos-*.xlsx")))
    if not arqs:
        return [], None
    ws = openpyxl.load_workbook(arqs[-1], data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Carteira")
    out = [dict(pagador=r[1], cnpj=dig(r[2]), nn=int(r[4]), rps=str(int(str(r[5]).strip())), emis=d(r[6]), venc=d(r[7]), valor=float(r[10]))
           for r in rows[h + 1:] if r and r[1]]
    return out, Path(arqs[-1]).name


# ------------------------------------------------------------------ montagem
def montar(cfg, D, hoje, saldos_omie):
    cli = {c["codigo_cliente_omie"]: (c.get("nome_fantasia") or c.get("razao_social") or "").strip() for c in D["CLI"]}
    wp = ler_webpilot()
    d0 = max(d(s["data"]) for s in cfg["saldos_documentados"])
    fim = d(cfg["fluxo_ate"])
    caixa0 = sum(s["valor"] for s in cfg["saldos_documentados"])

    # ---- títulos a receber
    tit = []
    for t in D["R"]:
        c, r = t["cabecTitulo"], t["resumo"]
        if c["cStatus"] == "CANCELADO":
            continue
        nf = str(c.get("cNumDocFiscal") or "")
        w = wp.get(nf, {})
        baixas = t.get("lancamentos", [])
        x = dict(id=c["nCodTitulo"], nf=nf, rps=rps_de(c.get("cNumTitulo")), emis=d(c["dDtEmissao"]), venc=d(c["dDtVenc"]),
                 prev=d(c.get("dDtPrevisao") or c["dDtVenc"]), status=c["cStatus"], liq=round(r["nValPago"] if c["cStatus"] == "RECEBIDO" else r["nValLiquido"] + r["nValAberto"], 2),
                 aberto=r["nValAberto"], cliente=w.get("cliente") or cli.get(c["nCodCliente"], "?"), navio=w.get("navio") or "—",
                 cnpj=dig(c.get("cCPFCNPJCliente")), receb=max((d(l["dDtLanc"]) for l in baixas), default=None))
        x["sit"] = "rec" if x["status"] == "RECEBIDO" else ("ab" if x["prev"] <= d0 else "av")
        tit.append(x)

    meses = {}
    for ym in cfg["meses_faturamento"]:
        a, m = map(int, ym.split("-"))
        lst = sorted([t for t in tit if t["emis"].year == a and t["emis"].month == m], key=lambda t: (t["emis"], t["nf"]))
        meses[ym] = lst

    # ---- fluxo
    dias = collections.defaultdict(lambda: {"ent": [], "sai": []})
    for t in tit:
        if t["status"] != "RECEBIDO" and d0 < t["prev"] <= fim:
            tag = "a confirmar" if t["prev"] < hoje else ""
            dias[t["prev"]]["ent"].append((t["cliente"] + (f" · {t['navio']}" if t["navio"] != "—" else "") + (f" <span class=\"pill pill-ab\" style=\"font-size:8px\">{tag}</span>" if tag else ""),
                                           f"NFSe {t['nf']}", t["aberto"]))
    est = cfg.get("estimativas", {})
    catn = cfg.get("categorias", {})
    for t in D["P"]:
        c = t["cabecTitulo"]
        if c["cStatus"] in ("PAGO", "LIQUIDADO", "CANCELADO"):
            continue
        v = d(c["dDtVenc"])
        if d0 < v <= fim:
            nota = est.get(str(c["nCodTitulo"]))
            dias[v]["sai"].append((esc(cli.get(c["nCodCliente"], "?")) + (f" <span class=\"pill pill-av\" style=\"font-size:8px\">estimativa</span>" if nota else ""),
                                   catn.get(c.get("cCodCateg") or "", "Tributos" if nota else "Outros"), t["resumo"]["nValAberto"] or c["nValorTitulo"]))
    for s in cfg.get("saidas_fora_do_omie", []):
        if d0 < d(s["data"]) <= fim:
            dias[d(s["data"])]["sai"].append((esc(s["descricao"]) + " <span class=\"pill pill-ab\" style=\"font-size:8px\">fora do Omie</span>", s["categoria"], s["valor"]))
    for x in D["L"]:  # realizado após o último extrato (ainda sem documento)
        dl = d(x["cabecalho"]["dDtLanc"])
        if dl > d0 and x["cabecalho"]["nCodCC"] in map(int, cfg["contas"]) and x["diversos"].get("cOrigem") not in ("BAXR", "BAXP", "TRAR", "TRAP"):
            k = "ent" if x["diversos"]["cNatureza"] == "R" else "sai"
            dias[dl][k].append((esc(x["detalhes"].get("cObs", ""))[:50] + " <span class=\"pill pill-av\" style=\"font-size:8px\">realizado s/ extrato</span>", "Omie", x["cabecalho"]["nValorLanc"]))

    prox = (hoje.replace(day=28) + dt.timedelta(days=5)).replace(day=1)
    cp_prox = []
    for t in D["P"]:
        c = t["cabecTitulo"]
        v = d(c["dDtVenc"])
        if c["cStatus"] not in ("PAGO", "LIQUIDADO", "CANCELADO") and (v.year, v.month) == (prox.year, prox.month):
            cp_prox.append(dict(dia=v.day, forn=cli.get(c["nCodCliente"], "?"), v=c["nValorTitulo"], cat=catn.get(c.get("cCodCateg") or "", "Outros"),
                                conta=cfg["contas"].get(str(c.get("nCodCC")), {}).get("nome", "—"), est=str(c["nCodTitulo"]) in est))

    fluxo, saldo = [], caixa0
    for dia in sorted(dias):
        e = sum(v for *_, v in dias[dia]["ent"]); s = sum(v for *_, v in dias[dia]["sai"])
        saldo += e - s
        fluxo.append(dict(dia=dia, ent=e, sai=s, saldo=saldo, itens=dias[dia]))

    # ---- cobrança: vencidos e boletos a baixar
    vencidos = sorted([t for t in tit if t["sit"] == "ab"], key=lambda t: -t["aberto"])
    bol, bol_arq = ler_boletos_banco()
    baixar, banco_por_tit = [], {}
    for b in bol:
        cand = [t for t in tit if t["rps"] == b["rps"] and abs((t["emis"] - b["emis"]).days) <= 30] or \
               [t for t in tit if t["cnpj"][:8] == b["cnpj"][:8] and abs(t["liq"] - b["valor"]) < 0.05 and abs((t["emis"] - b["emis"]).days) <= 30]
        if cand:
            t = min(cand, key=lambda t: abs((t["emis"] - b["emis"]).days))
            banco_por_tit[t["id"]] = b
            if t["status"] == "RECEBIDO":
                baixar.append((b, t))

    # ---- conciliação por conta (até o último extrato)
    concil = []
    for s in cfg["saldos_documentados"]:
        conta, ds = int(s["conta"]), d(s["data"])
        nc = [x for x in D["L"] if x["cabecalho"]["nCodCC"] == conta and d(x["cabecalho"]["dDtLanc"]) <= ds and not x["diversos"].get("dDtConc")]
        concil.append(dict(conta=cfg["contas"][s["conta"]]["detalhe"], data=ds, doc=s["documento"], banco=s["valor"],
                           omie=saldos_omie.get(s["conta"]), pend=len(nc), pend_v=sum(x["cabecalho"]["nValorLanc"] for x in nc)))

    # ---- histórico de saídas por categoria (6 meses fechados)
    ref = hoje.replace(day=1)
    hm = []
    for i in range(6, 0, -1):
        y, m = ref.year, ref.month - i
        while m <= 0:
            m += 12; y -= 1
        hm.append((y, m))
    hist = collections.defaultdict(lambda: collections.defaultdict(float))
    for x in D["L"]:
        if x["diversos"]["cNatureza"] != "P" or x["cabecalho"]["nCodCC"] not in map(int, cfg["contas"]):
            continue
        cat = x["detalhes"].get("cCodCateg") or ""
        if cat.startswith("0.01") or x["diversos"].get("cOrigem") in ("TRAP",):
            continue
        dl = d(x["cabecalho"]["dDtLanc"])
        if (dl.year, dl.month) in hm:
            hist[catn.get(cat, "Outros")][(dl.year, dl.month)] += x["cabecalho"]["nValorLanc"]

    return dict(cp_prox=cp_prox, prox=prox, tit=tit, d0=d0, fim=fim, caixa0=caixa0, meses=meses, fluxo=fluxo, vencidos=vencidos, banco=banco_por_tit,
                baixar=baixar, bol_arq=bol_arq, concil=concil, hist=hist, hm=hm, wp_meses=sorted({k for k in wp}))


# ------------------------------------------------------------------ html
def logo(tipo):
    b = base64.b64encode((RAIZ / "template" / f"logo-exact-{tipo}.png").read_bytes()).decode()
    return f'<img class="logo" src="data:image/png;base64,{b}" alt="Exact BR">'


def pagina(n, tot, eyebrow, sub, corpo, rodape, dark=False):
    return f"""<section class="page{' dark' if dark else ''}">
  <div class="page-header">{logo('escuro' if dark else 'claro')}<div class="meta"><span class="accent">{eyebrow}</span><span class="pipe"></span><span>{sub}</span></div></div>
  <div class="page-body">{corpo}</div>
  <div class="page-footer"><span>{rodape}</span><span class="page-num"><span class="cur">{n:02d}</span> / {tot:02d}</span></div>
</section>"""


def cards(itens):
    return '<div class="nums-grid" style="grid-template-columns:repeat(4,1fr)">' + "".join(
        f'<div class="num-cell"><div class="n-eye">{a}</div><div class="n-val {cor}">{v}</div><div class="n-sub">{s}</div></div>' for a, v, cor, s in itens) + "</div>"


def tabela_faturamento(lst):
    pill = {"rec": ("st-rec", "pill-rec", "Recebido"), "av": ("st-av", "pill-av", "A Vencer"), "ab": ("st-ab", "pill-ab", "Em Aberto")}
    linhas = "".join(
        f'<tr class="{pill[t["sit"]][0]}"><td>{dm(t["emis"])}</td><td>{esc(t["cliente"])}</td><td>{esc(t["navio"])}</td><td>{t["nf"][-4:]}</td>'
        f'<td>{dm(t["prev"])}</td><td>{dm(t["receb"])}</td><td class="num">{br(t["liq"])}</td><td><span class="pill {pill[t["sit"]][1]}">{pill[t["sit"]][2]}</span></td></tr>'
        for t in lst)
    return ('<table class="diag-table"><thead><tr><th>Emi</th><th>Cliente</th><th>Navio / Manobra</th><th>NFSe</th><th>Venc</th><th>Receb</th>'
            f'<th class="num">R$ Líquido</th><th>Status</th></tr></thead><tbody>{linhas}</tbody></table>')


def secao_mes(ym, lst, hoje, tem_wp):
    a, m = map(int, ym.split("-"))
    tot = sum(t["liq"] for t in lst)
    rec = [t for t in lst if t["sit"] == "rec"]; av = [t for t in lst if t["sit"] == "av"]; ab = [t for t in lst if t["sit"] == "ab"]
    parcial = (a, m) == (hoje.year, hoje.month)
    fonte = "Navio e manobra vindos do Webpilot." if tem_wp else "Sem exportação do Webpilot deste mês: navio fica em branco até a próxima carga."
    corpo = f"""<div class="chapter-num">— {{N}} —</div>
    <h2 class="editorial">O que foi faturado <span class="light">em</span> <span class="ital">{MESES[m-1]}.</span></h2>
    <p class="lede">{len(lst)} NFSes{' emitidas até ' + hoje.strftime('%d/%m') if parcial else ' emitidas'}. Status: <strong>Recebido</strong> (verde) · <strong>A Vencer</strong> (laranja) · <strong>Em Aberto</strong> = vencido sem baixa (vermelho). {fonte}</p>
    {cards([("Total faturado", "R$ " + br(tot, 0), "", f"{len(lst)} títulos · valor líquido"),
            ("Recebido", "R$ " + br(sum(t['liq'] for t in rec), 0), "green", f"{len(rec)} baixados"),
            ("A Vencer", "R$ " + br(sum(t['aberto'] for t in av), 0), "orange", f"{len(av)} títulos"),
            ("Em Aberto (vencido)", "R$ " + br(sum(t['aberto'] for t in ab), 0), "red", f"{len(ab)} títulos")])}
    <div class="diag-section"><h3>NFSes emitidas em {MESES[m-1]}</h3>
    <div class="sub">Valor líquido (após retenções) · Venc = previsão no Omie · Receb = data da baixa</div>{tabela_faturamento(lst)}</div>"""
    return corpo, tot, rec, av, ab


def secao_fluxo(M, cfg):
    sd = "".join(f'{cfg["contas"][s["conta"]]["nome"]} R$ {br(s["valor"])} ({s["documento"]})' + (" · " if i < len(cfg["saldos_documentados"]) - 1 else "")
                 for i, s in enumerate(cfg["saldos_documentados"]))
    rows = []
    for f in M["fluxo"]:
        sub = ""
        for k, tit, cor, sinal in (("ent", "Entradas", "#059669", "+"), ("sai", "Saídas", "#DC2626", "−")):
            if f["itens"][k]:
                linhas = "".join(f'<tr><td>{a}</td><td style="color:var(--gray-500);font-size:10px">{esc(b)}</td><td class="num" style="color:{cor}">{sinal} {br(v)}</td></tr>'
                                 for a, b, v in sorted(f["itens"][k], key=lambda x: -x[2]))
                sub += f'<div class="dia-sub"><div class="dia-sub-title">{tit}</div><table class="dia-sub-table">{linhas}</table></div>'
        var = f["ent"] - f["sai"]
        rows.append(f"""<details class="dia-row"><summary class="dia-summary"><span class="dia-chevron">▸</span>
    <span class="dia-label">{dm(f['dia'])} · {DIAS[f['dia'].weekday()]}</span>
    <span class="dia-ent num">{'+ ' + br(f['ent']) if f['ent'] else '—'}</span><span class="dia-sai num">{'− ' + br(f['sai']) if f['sai'] else '—'}</span>
    <span class="dia-delta num {'gap-neg' if var >= 0 else 'gap-pos'}">{'+' if var >= 0 else '−'} {br(abs(var))}</span><span class="dia-saldo num">{br(f['saldo'])}</span>
  </summary><div class="dia-content">{sub}</div></details>""")
    final = M["fluxo"][-1]["saldo"] if M["fluxo"] else M["caixa0"]
    corpo = f"""<div class="chapter-num">— {{N}} —</div>
    <h2 class="editorial">O caixa dia a dia <span class="light">até</span><br><span class="orange">{M['fim'].day:02d} de {MESES[M['fim'].month-1]}.</span></h2>
    <p class="lede">Clique em cada dia para abrir entradas e saídas. Saldo inicial = extratos de {M['d0'].strftime('%d/%m/%Y')}, únicos documentos de banco disponíveis: {sd}.</p>
    <div class="highlight-box"><div class="highlight-tag">Como a previsão foi montada</div>
    <p>Entradas: títulos em aberto no Omie pela data de previsão (boleto) — os marcados <strong>a confirmar</strong> já venceram depois do último extrato. Saídas: contas a pagar do Omie, que refletem os pagamentos recorrentes do histórico (seção seguinte), e as guias de imposto ainda <strong>fora do Omie</strong>. Fora do fluxo: {esc(cfg['fora_do_fluxo'])}.</p></div>
    <div class="fluxo-head"><div></div><div>Dia</div><div class="num">Entradas</div><div class="num">Saídas</div><div class="num">Variação</div><div class="num">Saldo projetado</div></div>
    <div class="saldo-inicial"><div></div><div>Saldo inicial · Itaú + Inter · {M['d0'].strftime('%d/%m/%Y')} (extratos)</div><div class="num" style="color:var(--orange);font-size:14px">R$ {br(M['caixa0'])}</div></div>
    {''.join(rows)}
    <div class="saldo-final"><div class="saldo-final-label">Saldo final projetado · {M['fim'].strftime('%d/%m/%Y')}</div><div class="num saldo-final-val">R$ {br(final)}</div></div>"""
    return corpo, final


def secao_cobranca(M):
    linhas_c = "".join(f'<tr><td>{esc(c["conta"])}</td><td>{dm(c["data"])}</td><td class="num">{br(c["banco"])}</td>'
                       f'<td class="num">{br(c["omie"]) if c["omie"] is not None else "—"}</td><td class="num">{br((c["omie"] or 0) - c["banco"])}</td><td>{c["pend"]}</td></tr>'
                       for c in M["concil"])
    def situ_banco(t):
        b = M["banco"].get(t["id"])
        return f"boleto vencido no Itaú (venc. {dm(b['venc'])})" if b else ("sem boleto em aberto no Itaú" if M["bol_arq"] else "—")
    linhas_v = "".join(f'<tr class="st-ab"><td>{esc(t["cliente"])}</td><td>{t["nf"]}</td><td>{dm(t["emis"])}</td><td>{dm(t["prev"])}</td>'
                       f'<td>{situ_banco(t)}</td><td class="num">{br(t["aberto"])}</td></tr>' for t in M["vencidos"])
    linhas_b = "".join(f'<tr><td>{esc(b["pagador"])[:28]}</td><td>{b["rps"]}</td><td>{b["nn"]}</td><td>{t["nf"]}</td><td>{dm(b["venc"])}</td>'
                       f'<td>{dm(t["receb"])}</td><td class="num">{br(b["valor"])}</td></tr>' for b, t in M["baixar"])
    tv = sum(t["aberto"] for t in M["vencidos"]); tb = sum(b["valor"] for b, _ in M["baixar"])
    corpo = f"""<div class="chapter-num">— {{N}} —</div>
    <h2 class="editorial">Conciliação <span class="light">e</span> <span class="orange">cobrança.</span></h2>
    <p class="lede">O que está conciliado contra o banco, o que está vencido de verdade e o que o banco ainda cobra sem precisar.</p>
    <div class="diag-section"><h3>A · Conciliação bancária até o último extrato</h3><div class="sub">Saldo do Omie na data do extrato contra o documento do banco</div>
    <table class="diag-table"><thead><tr><th>Conta</th><th>Extrato</th><th class="num">Banco</th><th class="num">Omie</th><th class="num">Diferença</th><th>Não conciliados</th></tr></thead><tbody>{linhas_c}</tbody></table></div>
    <div class="diag-section"><h3>B · Vencidos em cobrança · R$ {br(tv)}</h3><div class="sub">Sem pagamento identificado em nenhum extrato, boleto ou TED</div>
    <table class="diag-table"><thead><tr><th>Cliente</th><th>NFSe</th><th>Emissão</th><th>Venc</th><th>No banco</th><th class="num">Em aberto</th></tr></thead><tbody>{linhas_v}</tbody></table></div>
    {f'''<div class="diag-section"><h3>C · Boletos a baixar no Itaú · R$ {br(tb)}</h3><div class="sub">Já recebidos no Omie (por TED), mas ainda vencidos no banco ({esc(M["bol_arq"])}) — risco de pagamento em dobro e tarifa</div>
    <table class="diag-table"><thead><tr><th>Pagador</th><th>RPS</th><th>Nosso nº</th><th>NFSe</th><th>Venc</th><th>Recebido</th><th class="num">Valor</th></tr></thead><tbody>{linhas_b}</tbody></table></div>''' if M["baixar"] else ''}"""
    return corpo, tv, tb


def secao_historico(M):
    cab = "".join(f"<th class=\"num\">{MESES[m-1][:3]}/{str(y)[2:]}</th>" for y, m in M["hm"])
    ordem = sorted(M["hist"], key=lambda k: -sum(M["hist"][k].values()))
    linhas = "".join(f'<tr><td>{esc(k)}</td>' + "".join(f'<td class="num">{br(M["hist"][k].get(ym, 0), 0) if M["hist"][k].get(ym) else "—"}</td>' for ym in M["hm"])
                     + f'<td class="num"><strong>{br(sum(M["hist"][k].values()) / len(M["hm"]), 0)}</strong></td></tr>' for k in ordem)
    return f"""<div class="chapter-num">— {{N}} —</div>
    <h2 class="editorial">O histórico <span class="light">que sustenta</span> <span class="orange">a previsão.</span></h2>
    <p class="lede">Saídas realizadas por categoria nos últimos seis meses fechados (Itaú + Inter, sem transferências entre contas). É daqui que saem os valores recorrentes do fluxo e os alertas de itens variáveis.</p>
    <div class="diag-section"><h3>Saídas por categoria · R$</h3><div class="sub">Média = total dos seis meses ÷ 6</div>
    <table class="diag-table"><thead><tr><th>Categoria</th>{cab}<th class="num">Média</th></tr></thead><tbody>{linhas}</tbody></table></div>
    <div class="highlight-box"><div class="highlight-tag">Leitura</div><p><strong>Fixos</strong> (pró-labore, contabilidade, sistemas, plano de saúde, conselho) já estão cadastrados como recorrência no Omie. <strong>ISS</strong> e <strong>IRPJ/CSLL</strong> variam com o faturamento — no fluxo aparecem marcados como estimativa até a guia sair. <strong>Distribuição de lucros</strong> e o IRRF sobre ela dependem da decisão dos sócios e ficam fora da projeção.</p></div>"""


def gerar(cfg, M, hoje, saida):
    sig, emp = cfg["sigla"], cfg["empresa"]
    yms = cfg["meses_faturamento"]
    nomes = " + ".join(MESES[int(y.split('-')[1]) - 1].capitalize() for y in yms)
    ano = yms[-1][:4]
    rod = f"Relatório · {sig} · {nomes.replace(' + ', '+')} {ano}"
    TOT = 8
    secs, totais = [], {}
    for ym in yms:
        a, m = map(int, ym.split("-"))
        tem_wp = any(t["navio"] != "—" for t in M["meses"][ym])
        corpo, tot, rec, av, ab = secao_mes(ym, M["meses"][ym], hoje, tem_wp)
        totais[ym] = (tot, rec, av, ab)
        secs.append((f"Faturamento {MESES[m-1].capitalize()}", corpo))
    fl, final = secao_fluxo(M, cfg); secs.append(("Fluxo de Caixa Projetado", fl))
    cb, tv, tb = secao_cobranca(M); secs.append(("Conciliação e Cobrança", cb))
    secs.append(("Histórico de Pagamentos", secao_historico(M)))

    ym1, ym2 = yms[0], yms[-1]
    t1, t2 = totais[ym1], totais[ym2]
    entra = sum(f["ent"] for f in M["fluxo"]); sai = sum(f["sai"] for f in M["fluxo"])
    hoje_sai = next((f["sai"] for f in M["fluxo"] if f["dia"] == hoje), 0)
    sumario = f"""<div class="chapter-num">— 01 —</div>
    <h2 class="editorial">O caixa de {MESES[hoje.month-1]} <span class="light">em</span><br><span class="orange">quatro vetores.</span></h2>
    <p class="lede">Faturamento dos dois meses, caixa documentado nos extratos e para onde ele vai até {M['fim'].strftime('%d/%m')}.</p>
    {cards([(f"Faturamento {MESES[int(ym1[5:])-1][:3].capitalize()}", mm(t1[0]), "orange", f"{len(M['meses'][ym1])} títulos · R$ {br(sum(t['liq'] for t in t1[1]) / 1e3, 0)}K recebidos"),
            (f"Faturamento {MESES[int(ym2[5:])-1][:3].capitalize()}", mm(t2[0]), "orange", f"{len(M['meses'][ym2])} títulos · até {hoje.strftime('%d/%m')}"),
            ("Caixa documentado", mm(M['caixa0']), "green", f"Itaú + Inter · extratos {M['d0'].strftime('%d/%m')}"),
            ("Vencidos em cobrança", mm(tv), "red", f"{len(M['vencidos'])} títulos sem pagamento")])}
    <h3>Leitura em quatro linhas</h3><ul class="editorial-list">
      <li><strong>{MESES[int(ym1[5:])-1].capitalize()} fechado.</strong> {len(M['meses'][ym1])} NFSes somando R$ {br(t1[0])} — R$ {br(sum(t['liq'] for t in t1[1]))} recebidos, R$ {br(sum(t['aberto'] for t in t1[3]))} vencidos e R$ {br(sum(t['aberto'] for t in t1[2]))} a vencer.</li>
      <li><strong>{MESES[int(ym2[5:])-1].capitalize()} em curso.</strong> {len(M['meses'][ym2])} títulos até {hoje.strftime('%d/%m')} somando R$ {br(t2[0])}; R$ {br(sum(t['aberto'] for t in t2[2]))} ainda a vencer.</li>
      <li><strong>Caixa até {M['fim'].strftime('%d/%m')}.</strong> Parte de R$ {br(M['caixa0'])} (extratos de {M['d0'].strftime('%d/%m')}), entram R$ {br(entra)} e saem R$ {br(sai)}{f' — R$ {br(hoje_sai)} vencem hoje' if hoje_sai else ''}; saldo projetado R$ {br(final)}.</li>
      <li><strong>Cobrança.</strong> R$ {br(tv)} vencidos sem pagamento identificado{f' e R$ {br(tb)} em boletos já pagos por TED que o Itaú ainda cobra' if tb else ''}.</li>
    </ul>"""

    paginas = [pagina(2, TOT, "Sumário Executivo", f"Posição {sig}", sumario, rod)]
    for i, (titulo, corpo) in enumerate(secs, start=2):
        paginas.append(pagina(i + 1, TOT, f"Seção {i:02d}", titulo, corpo.replace("{N}", f"{i:02d}"), rod))

    passos = []
    if hoje_sai:
        passos.append(f"<strong style=\"color:#FFF\">Hoje.</strong> Aprovar R$ {br(hoje_sai)} em pagamentos que vencem hoje (guias de imposto, pró-labore e contabilidade).")
    passos.append(f"<strong style=\"color:#FFF\">Esta semana.</strong> Enviar extratos Itaú e Inter atualizados para confirmar os boletos vencidos após {M['d0'].strftime('%d/%m')} e cadastrar no Omie as guias que estão fora dele.")
    if tv:
        nb = len(M["baixar"])
        extra = f" e pedir ao Itaú a baixa de {nb} boletos já pagos por TED" if nb else ""
        passos.append(f"<strong style=\"color:#FFF\">Cobrança.</strong> Cobrar R$ {br(tv)} vencidos{extra}.")
    if cfg.get("estimativas"):
        passos.append("<strong style=\"color:#FFF\">Antes do dia 30.</strong> Receber da JFS a guia de IRPJ/CSLL do trimestre — hoje o fluxo usa estimativa.")
    li = "".join(f'<li style="color:rgba(255,255,255,0.85);border-bottom-color:rgba(255,255,255,0.12)">{p}</li>' for p in passos)
    capa = f"""<section class="page dark"><div class="page-header">{logo('escuro')}<div class="meta"><span class="accent">BPO Mensal</span><span class="pipe"></span><span>{MESES[hoje.month-1].capitalize()} · {hoje.year}</span></div></div>
  <div class="cover-content"><div class="cover-eyebrow">{esc(cfg['para'])}</div>
    <h1 class="cover-title">Posição mensal<br><span class="ital">de {nomes.lower().replace(' + ', ' e ')}.</span></h1>
    <p class="cover-sub">Faturamento com navio e status de recebimento, previsão de caixa clicável construída sobre o histórico de pagamentos e a conciliação de títulos e cobranças — tudo em uma leitura só.</p>
    <div class="cover-meta"><div><div class="label">Cliente</div><div class="value">{esc(emp)}</div></div>
      <div><div class="label">Período</div><div class="value ital">{nomes} / {ano}</div></div>
      <div><div class="label">Data</div><div class="value">{hoje.day} de {MESES[hoje.month-1]} de {hoje.year}</div></div>
      <div><div class="label">Por</div><div class="value">João Victor<br>Exact BR</div></div></div></div></section>"""
    fecho = f"""<section class="page dark"><div class="page-header">{logo('escuro')}<div class="meta"><span class="accent">Próximo Passo</span></div></div>
  <div class="page-body centered"><div class="chapter-num">— final —</div>
    <h2 class="editorial">O caixa de {MESES[M['fim'].month-1]}<br><span class="light">começa com</span> <span class="orange">documento na mão.</span></h2>
    <ul class="editorial-list" style="margin-top:24px">{li}</ul>
    <div class="signature-grid"><div class="signature"><div class="sig-name">João Victor Andrade de Souza</div><div class="sig-role">Sócio-fundador · Exact BR</div></div>
      <div style="text-align:right"><div class="sig-role" style="margin-bottom:4px">Disciplina Gera Lucro</div><div style="font-family:'Plus Jakarta Sans',sans-serif;font-size:13px;font-weight:500;color:rgba(255,255,255,0.85)">Exact BR · Recife — PE</div></div></div></div></section>"""
    head = (RAIZ / "template" / "head.html").read_text(encoding="utf-8").replace("{{TITULO}}", f"Posição {sig} · {nomes.upper().replace(' + ', '+')} {ano} · Exact BR")
    saida.parent.mkdir(parents=True, exist_ok=True)
    saida.write_text(head + "<body>\n" + capa + "\n".join(paginas) + fecho + '\n<button class="print-btn" onclick="window.print()">↓ Exportar PDF</button>\n</body>\n</html>\n', encoding="utf-8")
    return final


def gerar_painel(cfg, M, hoje, saida):
    st = {"rec": "Recebido", "av": "A Vencer", "ab": "Em Aberto"}
    curto = lambda nf: int(nf[-4:]) if nf and nf[-4:].isdigit() else None
    def t2d(t):
        return dict(doc=t["rps"], nf=curto(t["nf"]), emi=t["emis"].isoformat(), venc=t["prev"].isoformat(), baixa=t["receb"].isoformat() if t["receb"] else None,
                    cli=t["cliente"], navio=t["navio"], liq=t["liq"], saldo=t["aberto"] if t["sit"] != "rec" else 0, st=st[t["sit"]])
    ids = set()
    tit = []
    for ym in cfg["meses_faturamento"]:
        for t in M["meses"][ym]:
            ids.add(t["id"]); tit.append(t2d(t))
    tit += [t2d(t) for t in M["vencidos"] if t["id"] not in ids]  # vencidos de meses anteriores entram na cobrança
    limpa = lambda h: re.sub(r"<[^>]+>", "", re.sub(r'<span class="pill[^"]*"[^>]*>[^<]*</span>', "", h)).strip()
    tag = lambda h: " · ".join(re.findall(r'<span class="pill[^"]*"[^>]*>([^<]+)</span>', h))
    dias = [dict(iso=f["dia"].isoformat(), lbl=dm(f["dia"]), dow=DIAS[f["dia"].weekday()].lower(), we=f["dia"].weekday() >= 5,
                 ent=[dict(desc=limpa(a), sub=" · ".join(x for x in (b, tag(a)) if x), v=v) for a, b, v in f["itens"]["ent"]],
                 sai=[dict(desc=limpa(a), sub=" · ".join(x for x in (b, tag(a)) if x), v=v) for a, b, v in f["itens"]["sai"]]) for f in M["fluxo"]]
    agg = collections.OrderedDict()
    for r in sorted(M["cp_prox"], key=lambda r: (r["dia"], r["forn"])):
        k = (r["dia"], r["forn"], r["cat"])
        a = agg.setdefault(k, dict(dia=r["dia"], forn=r["forn"], v=0, n=0, conta=r["conta"], obs=r["cat"], tag="wait" if r["est"] else ""))
        a["v"] += r["v"]; a["n"] += 1
    rec = [dict(dia=a["dia"], forn=a["forn"], val=br(a["v"]) + (f" ({a['n']}×)" if a["n"] > 1 else ""), obs=a["obs"] + (" · valor estimado" if a["tag"] else ""), conta=a["conta"], tag=a["tag"]) for a in agg.values()]
    rec = sorted(rec + cfg.get("recorrentes_variaveis", []), key=lambda r: r["dia"])
    contas = [dict(nome=cfg["contas"][s["conta"]]["nome"], cc=cfg["contas"][s["conta"]].get("cc", ""), obs=cfg["contas"][s["conta"]].get("obs", ""), v=s["valor"]) for s in cfg["saldos_documentados"]]
    yms = cfg["meses_faturamento"]
    meses = [[ym, MESES[int(ym[5:]) - 1][:3].upper(), MESES[int(ym[5:]) - 1].capitalize()] for ym in yms]
    dados = dict(caixa=M["caixa0"], contas=contas, d0lbl=dm(M["d0"]), fimlbl=dm(M["fim"]), hoje=hoje.isoformat(), hojelbl=dm(hoje), hojeDia=hoje.day,
                 meses=meses, tit=tit, days=dias, rec=rec,
                 baixar=[dict(venc=b["venc"].isoformat(), pagador=b["pagador"], rps=b["rps"], nn=b["nn"], nf=t["nf"][-4:], receb=t["receb"].isoformat() if t["receb"] else None, valor=b["valor"]) for b, t in M["baixar"]],
                 concil=[dict(conta=c["conta"], data=c["data"].isoformat(), banco=c["banco"], omie=c["omie"] or 0, pend=c["pend"]) for c in M["concil"]])
    chips = '<button class="chip" data-f="m" data-v="" aria-pressed="true">' + " + ".join(m[1] for m in meses) + '</button>' + "".join(
        f'<button class="chip" data-f="m" data-v="{m[0]}" aria-pressed="false">{m[1]}</button>' for m in meses)
    docs = " e ".join(f'{cfg["contas"][s["conta"]]["nome"]} {dm(d(s["data"]))}' for s in cfg["saldos_documentados"])
    rod = (f'<span><b>Fontes:</b> Omie (títulos, baixas e contas a pagar), Webpilot (navio), extratos {docs}, relatório de boletos do Itaú.</span>'
           f'<span><b>No fluxo:</b> títulos a receber pela previsão, contas a pagar do Omie e guias cadastradas. <b>Fora:</b> {esc(cfg["fora_do_fluxo"])}.</span>'
           '<span>Exact BR · Disciplina Gera Lucro</span>')
    tpl = (RAIZ / "template" / "dashboard.html").read_text(encoding="utf-8")
    rep = {"{{DADOS}}": json.dumps(dados, ensure_ascii=False), "{{TITULO}}": f"Painel {cfg['sigla']} · {cfg['empresa']}", "{{EMPRESA}}": esc(cfg["empresa"]),
           "{{POSICAO}}": f"Posição {hoje:%d/%m/%Y} · extratos {M['d0']:%d/%m}", "{{JANELA}}": f"{dm(M['d0'] + dt.timedelta(days=1))} → {dm(M['fim'])}",
           "{{FATHINT}}": f"Omie + Webpilot · {len(tit)} títulos · valor líquido · Em Aberto = vencido (inclui vencidos anteriores)",
           "{{CHIPSMES}}": chips, "{{BOLARQ}}": esc(M["bol_arq"] or "sem relatório do banco"), "{{RODAPE}}": rod}
    for k, v in rep.items():
        tpl = tpl.replace(k, v)
    saida.write_text(tpl, encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default=str(RAIZ / "config" / "bmg.json"))
    ap.add_argument("--hoje", default=dt.date.today().isoformat())
    ap.add_argument("--offline", action="store_true")
    a = ap.parse_args()
    cfg = json.loads(Path(a.config).read_text(encoding="utf-8"))
    hoje = dt.date.fromisoformat(a.hoje)
    cache = RAIZ / "dados" / "omie.json"
    D = json.loads(cache.read_text()) if a.offline else buscar_omie(cache)
    so_cache = RAIZ / "dados" / "saldos_omie.json"
    if a.offline and so_cache.exists():
        so = json.loads(so_cache.read_text())
    else:
        so = {s["conta"]: saldo_omie(s["conta"], d(s["data"])) for s in cfg["saldos_documentados"]}
        so_cache.write_text(json.dumps(so))
    M = montar(cfg, D, hoje, so)
    yms = cfg["meses_faturamento"]
    saida = RAIZ / "relatorios" / hoje.strftime("%Y-%m") / f"posicao-{cfg['sigla'].lower()}-{'-'.join(MESES[int(y[5:])-1][:3] for y in yms)}-{yms[-1][:4]}.html"
    final = gerar(cfg, M, hoje, saida)
    (RAIZ / "index.html").write_text(saida.read_text(encoding="utf-8"), encoding="utf-8")  # última posição publicada na raiz (GitHub Pages)
    painel = saida.with_name(saida.name.replace("posicao-", "painel-"))
    gerar_painel(cfg, M, hoje, painel)
    (RAIZ / "painel.html").write_text(painel.read_text(encoding="utf-8"), encoding="utf-8")
    print(f"ok → {painel.relative_to(RAIZ)} e painel.html")
    print(f"ok → {saida.relative_to(RAIZ)} | caixa {br(M['caixa0'])} → {br(final)} em {M['fim']:%d/%m} | vencidos {len(M['vencidos'])} | boletos a baixar {len(M['baixar'])}")


if __name__ == "__main__":
    main()
